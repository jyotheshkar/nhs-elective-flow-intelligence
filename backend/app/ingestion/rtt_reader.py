"""Reader for NHS England provider-level RTT workbooks.

The published workbooks are ~10 MB spreadsheets with roughly 130 columns, of
which this product uses about a dozen. The reader therefore streams the sheet
in openpyxl read-only mode and materialises only the mapped columns.

Layout of every provider workbook inspected (Apr-2023 onwards):

    rows 1-11   publication metadata (Period, Source, Published, Status ...)
    row  12     section banner
    row  13     column headers
    row  14+    one row per provider x treatment function

The header row is configured rather than hardcoded, and is re-detected if the
configured position does not contain the expected captions.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from app.config import get_source_config
from app.transformation.schema import (
    DATASET_SPECS,
    DatasetSpec,
    build_column_map,
    normalise_caption,
)
from app.utils.logging import get_logger

logger = get_logger(__name__)

_METADATA_KEYS = {
    "period:": "source_period",
    "source:": "source_system",
    "published:": "published_date",
    "revised:": "revised",
    "status:": "status",
    "basis:": "basis",
}


class WorkbookSchemaError(RuntimeError):
    """Raised when a workbook does not match the expected provider schema."""


@dataclass
class ParsedWorkbook:
    """A parsed workbook: canonical rows plus the publication metadata."""

    dataset: str
    frame: pd.DataFrame
    metadata: dict[str, str]
    header_row: int
    mapped_columns: dict[str, str]
    unmapped_required: list[str]


def _find_sheet(workbook: Any, preferred: str) -> Any:
    """Return the provider sheet, tolerating small naming changes."""
    if preferred in workbook.sheetnames:
        return workbook[preferred]
    for name in workbook.sheetnames:
        lowered = name.strip().lower()
        # "IS Provider" is independent-sector-only and must not be picked up.
        if lowered == "provider":
            return workbook[name]
    raise WorkbookSchemaError(
        f"No '{preferred}' sheet found. Available sheets: {workbook.sheetnames}"
    )


def _read_metadata(rows: list[list[Any]]) -> dict[str, str]:
    """Extract the publication metadata block from the first rows."""
    metadata: dict[str, str] = {}
    for row in rows:
        for index, cell in enumerate(row[:4]):
            key = normalise_caption(cell)
            if key in _METADATA_KEYS and index + 1 < len(row):
                value = row[index + 1]
                if value is not None and str(value).strip():
                    metadata[_METADATA_KEYS[key]] = str(value).strip()
    return metadata


def _locate_header(rows: list[list[Any]], configured: int) -> int:
    """Return the 0-indexed header row, preferring the configured position.

    Falls back to scanning for the row that contains 'Provider Code', so a
    layout shift in a future publication degrades to a warning rather than a
    parse failure.
    """
    def looks_like_header(row: list[Any]) -> bool:
        captions = {normalise_caption(c) for c in row}
        return "provider code" in captions and "treatment function code" in captions

    configured_index = configured - 1
    if 0 <= configured_index < len(rows) and looks_like_header(rows[configured_index]):
        return configured_index
    for index, row in enumerate(rows):
        if looks_like_header(row):
            logger.warning(
                "Header row moved: expected row %d, found row %d",
                configured, index + 1,
            )
            return index
    raise WorkbookSchemaError(
        "Could not locate a header row containing 'Provider Code' and "
        "'Treatment Function Code'. The workbook layout has changed."
    )


def _coerce_numeric(value: Any) -> float | None:
    """Parse a published cell into a number.

    NHS England uses '-', '*' (small-number suppression) and blanks. All become
    null rather than zero: treating a suppressed value as zero would understate
    waiting lists and silently bias every derived metric.
    """
    if value is None:
        return None
    if isinstance(value, int | float) and not isinstance(value, bool):
        return float(value)
    text = str(value).strip().replace(",", "").replace("%", "")
    if text in {"", "-", "*", "n/a", "N/A", ".", "..", "z", "c"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def read_provider_workbook(
    path: Path,
    dataset: str,
    reporting_month: date,
) -> ParsedWorkbook:
    """Parse one provider-level workbook into canonical columns."""
    spec: DatasetSpec = DATASET_SPECS[dataset]
    rtt_cfg = get_source_config()["sources"]["rtt"]
    sheet_name = rtt_cfg["datasets"][dataset].get("sheet", "Provider")
    configured_header = int(rtt_cfg.get("header_row", 13))

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = _find_sheet(workbook, sheet_name)
        row_iter = sheet.iter_rows(values_only=True)

        # Buffer the preamble so metadata and header detection can look around.
        preamble: list[list[Any]] = []
        header_index: int | None = None
        for row in row_iter:
            preamble.append(list(row))
            if len(preamble) >= max(configured_header, 1):
                try:
                    header_index = _locate_header(preamble, configured_header)
                except WorkbookSchemaError:
                    if len(preamble) > 40:
                        raise
                    continue
                break
        if header_index is None:
            header_index = _locate_header(preamble, configured_header)

        metadata = _read_metadata(preamble[:header_index])
        captions = preamble[header_index]
        column_map = build_column_map(captions, spec)

        mapped = set(column_map.values())
        unmapped_required = [m for m in spec.required_measures if m not in mapped]
        missing_dimensions = [
            d for d in ("provider_code", "specialty_code") if d not in mapped
        ]
        if missing_dimensions:
            raise WorkbookSchemaError(
                f"Workbook {path.name} is missing required dimension columns: "
                f"{missing_dimensions}"
            )
        if unmapped_required:
            raise WorkbookSchemaError(
                f"Workbook {path.name} is missing required measures: "
                f"{unmapped_required}. Published captions may have changed - "
                "update backend/app/transformation/schema.py."
            )

        text_fields = {
            "region_code", "provider_code", "provider_name",
            "specialty_code", "specialty_name",
        }
        indices = sorted(column_map)
        records: list[dict[str, Any]] = []
        # Rows already consumed into the preamble but sitting after the header.
        for row in preamble[header_index + 1:]:
            record = _row_to_record(row, column_map, indices, text_fields)
            if record:
                records.append(record)
        for row in row_iter:
            record = _row_to_record(list(row), column_map, indices, text_fields)
            if record:
                records.append(record)
    finally:
        workbook.close()

    frame = pd.DataFrame.from_records(records)
    if frame.empty:
        raise WorkbookSchemaError(f"No data rows parsed from {path.name}.")

    frame["reporting_month"] = pd.Timestamp(reporting_month)
    logger.info(
        "Parsed %s: %d rows, %d canonical columns",
        path.name, len(frame), len(column_map),
    )
    return ParsedWorkbook(
        dataset=dataset,
        frame=frame,
        metadata=metadata,
        header_row=header_index + 1,
        mapped_columns={str(i): name for i, name in column_map.items()},
        unmapped_required=unmapped_required,
    )


def _row_to_record(
    row: list[Any],
    column_map: dict[int, str],
    indices: list[int],
    text_fields: set[str],
) -> dict[str, Any] | None:
    """Convert one worksheet row into a canonical record, or None to skip it."""
    if not row:
        return None
    record: dict[str, Any] = {}
    for index in indices:
        if index >= len(row):
            continue
        name = column_map[index]
        value = row[index]
        if name in text_fields:
            record[name] = None if value is None else str(value).strip()
        else:
            record[name] = _coerce_numeric(value)
    # A row without a provider code is a spacer, footnote or grand-total band.
    if not record.get("provider_code"):
        return None
    return record
